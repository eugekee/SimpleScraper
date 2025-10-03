import pandas as pd
from datetime import datetime
import os
import logging
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference

logger = logging.getLogger(__name__)


def save_to_excel(data, filename=None, sheet_name=None):
    try:
        if filename is None:
            from config import EXCEL_FILE, EXCEL_SHEET
            filename = EXCEL_FILE
            sheet_name = EXCEL_SHEET or "Data"

        os.makedirs(os.path.dirname(os.path.abspath(filename)), exist_ok=True)
        new_df = pd.DataFrame([data])

        if os.path.exists(filename):
            try:
                existing_df = pd.read_excel(filename, sheet_name=sheet_name, engine='openpyxl')
                updated_df = pd.concat([existing_df, new_df], ignore_index=True)
            except Exception as e:
                logger.warning("Ошибка при чтении существующего файла, создаем новый: %s", str(e))
                updated_df = new_df
        else:
            updated_df = new_df

        updated_df.to_excel(filename, sheet_name=sheet_name, index=False, engine='openpyxl')

        create_excel_chart(filename, sheet_name)

        logger.info("Данные успешно сохранены в %s", filename)
        return True

    except Exception as e:
        logger.error("Ошибка при сохранении в Excel: %s", str(e))
        try:
            csv_filename = str(filename).replace('.xlsx', '.csv')
            if os.path.exists(csv_filename):
                new_df.to_csv(csv_filename, mode='a', header=False, index=False)
            else:
                new_df.to_csv(csv_filename, index=False)
            logger.info("Данные сохранены в CSV: %s", csv_filename)
            return True
        except Exception as csv_error:
            logger.error("Ошибка при сохранении в CSV: %s", str(csv_error))
            return False


def create_excel_chart(filename, sheet_name):

    try:
        workbook = load_workbook(filename)
        sheet = workbook[sheet_name]

        df = pd.read_excel(filename, sheet_name=sheet_name)

        if len(df) < 2:
            logger.info("Недостаточно данных для построения графика")
            return

        numeric_data = []
        for index, row in df.iterrows():
            try:
                price = float(str(row['price']).replace(',', '.'))

                if 100 <= price <= 500:
                    numeric_data.append({
                        'timestamp': row['timestamp'],
                        'price': price
                    })
            except (ValueError, TypeError):
                continue

        if len(numeric_data) < 2:
            logger.info("Недостаточно числовых данных для построения графика")
            return

        filtered_df = pd.DataFrame(numeric_data)

        for chart in sheet._charts:
            sheet._charts.remove(chart)

        chart = LineChart()
        chart.title = "Динамика цены акций Сбербанка"
        chart.style = 13
        chart.x_axis.title = "Время"
        chart.y_axis.title = "Цена (руб.)"
        chart.height = 15
        chart.width = 30

        data_len = len(filtered_df) + 1
        data = Reference(sheet, min_col=2, min_row=1, max_col=2, max_row=data_len)
        categories = Reference(sheet, min_col=1, min_row=2, max_row=data_len)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        chart.legend = None
        sheet.add_chart(chart, f"G2")
        workbook.save(filename)
        logger.info("График успешно добавлен в Excel файл")

    except Exception as e:
        logger.error("Ошибка при создании графика: %s", str(e))


def read_excel_data(filename=None, sheet_name=None):
    try:
        if filename is None:
            from config import EXCEL_FILE, EXCEL_SHEET
            filename = EXCEL_FILE
            sheet_name = EXCEL_SHEET or "Data"

        if os.path.exists(filename):
            df = pd.read_excel(filename, sheet_name=sheet_name, engine='openpyxl')
            return df
        else:
            logger.warning("Файл %s не существует", filename)
            return pd.DataFrame()
    except Exception as e:
        logger.error("Ошибка при чтении Excel: %s", str(e))
        return pd.DataFrame()

def get_sber_price_fallback():
    try:
        import random
        base_price = 300.0
        price = round(base_price + random.uniform(-10, 10), 2)
        change = round(random.uniform(-5, 5), 2)
        change_percent = round((change / base_price) * 100, 2)

        return {
            'timestamp': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'price': str(price),
            'change': f"{'+' if change >= 0 else ''}{change}",
            'change_percent': f"{'+' if change_percent >= 0 else ''}{change_percent}",
            'high': str(round(price + random.uniform(1, 5), 2)),
            'low': str(round(price - random.uniform(1, 5), 2)),
            'volume': str(random.randint(1000000, 5000000)),
            'source': 'fallback_test'
        }
    except Exception as e:
        logger.error("Fallback метод не сработал: %s", str(e))
        return None


def monitor_sber_stock():
    data = None
    sources_tried = []

    try:
        try:
            from alternative_scraper import get_sber_price_simple
            logger.info("Попытка получить данные через простой скрапер...")
            data = get_sber_price_simple()
            if data and data.get('price') not in [None, 'N/A']:
                logger.info("Данные получены через простой скрапер")
                data['source'] = 'simple_scraper'
                sources_tried.append('simple_scraper: success')
            else:
                data = None
                sources_tried.append('simple_scraper: failed')
        except Exception as e:
            logger.warning("Ошибка в простом скрапере: %s", str(e))
            sources_tried.append('simple_scraper: error')
        if data is None:
            try:
                from scraper import scrape_sber_stock
                logger.info("Попытка получить данные через основной скрапер...")
                data = scrape_sber_stock()
                if data and data.get('price') not in [None, 'N/A', 'SCRAPER_ERROR', 'PARSE_ERROR']:
                    logger.info("Данные получены через основной скрапер")
                    data['source'] = 'main_scraper'
                    sources_tried.append('main_scraper: success')
                else:
                    data = None
                    sources_tried.append('main_scraper: failed')
            except Exception as e:
                logger.warning("Ошибка в основном скрапере: %s", str(e))
                sources_tried.append('main_scraper: error')
        if data is None:
            logger.info("Все методы не сработали, используем тестовые данные...")
            data = get_sber_price_fallback()
            if data:
                logger.info("Используются тестовые данные")
                sources_tried.append('fallback: success')
            else:
                sources_tried.append('fallback: failed')
        if isinstance(data, dict):
            for key in ['price', 'change', 'change_percent', 'high', 'low', 'volume']:
                if data.get(key) in [None, '']:
                    data[key] = 'N/A'
            if 'timestamp' not in data:
                data['timestamp'] = datetime.now().strftime('%Y-%m-%d %H:%M:%S')

            if 'source' not in data:
                data['source'] = 'unknown'

            save_success = save_to_excel(data)
            if save_success:
                logger.info("Данные успешно сохранены. Цена: %s, Источник: %s", data.get('price'), data.get('source'))
            else:
                logger.warning("Данные получены, но не сохранены")

            return data
        else:
            logger.error("Получены некорректные данные типа: %s", type(data))
            error_data = {
                'timestamp': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
                'price': 'PARSE_ERROR',
                'change': 'PARSE_ERROR',
                'change_percent': 'PARSE_ERROR',
                'high': 'PARSE_ERROR',
                'low': 'PARSE_ERROR',
                'volume': 'PARSE_ERROR',
                'source': 'parse_error'
            }
            save_to_excel(error_data)
            return error_data

    except Exception as e:
        logger.error("Критическая ошибка в мониторинге: %s", str(e))
        error_data = {
            'timestamp': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'price': 'MONITOR_ERROR',
            'change': 'MONITOR_ERROR',
            'change_percent': 'MONITOR_ERROR',
            'high': 'MONITOR_ERROR',
            'low': 'MONITOR_ERROR',
            'volume': 'MONITOR_ERROR',
            'source': 'monitor_error'
        }
        try:
            save_to_excel(error_data)
        except Exception as save_error:
            logger.error("Ошибка при сохранении данных об ошибке: %s", str(save_error))
        return error_data